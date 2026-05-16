import openpyxl
import os
import json
from datetime import datetime

def audit_pics(public_path, engineering_path):
    results = {
        "metadata": {
            "date": datetime.now().isoformat(),
            "public_file": os.path.basename(public_path),
            "engineering_file": os.path.basename(engineering_path)
        },
        "models": {}
    }

    pub_wb = openpyxl.load_workbook(public_path, data_only=True)
    eng_wb = openpyxl.load_workbook(engineering_path, data_only=True)

    # All sheets that look like SunSpec models (purely numeric or known models)
    model_sheets = [s for s in pub_wb.sheetnames if s.isdigit() or s in ['Certification Info', 'Testing Info']]
    
    # We also want to check for sheets present in Eng but NOT in Pub (the Delta)
    eng_only_sheets = [s for s in eng_wb.sheetnames if s not in pub_wb.sheetnames]
    results["extensions"] = eng_only_sheets

    for model_id in model_sheets:
        if not model_id.isdigit(): continue
        
        results["models"][model_id] = {
            "status": "Present in both",
            "points": {}
        }
        
        pub_sheet = pub_wb[model_id]
        eng_sheet = eng_wb[model_id] if model_id in eng_wb.sheetnames else None
        
        if not eng_sheet:
            results["models"][model_id]["status"] = "Present in Public ONLY"
            continue

        # Extract points from sheets
        # Assumption: PICS format has point names in a specific column. 
        # Typically Col B or C. Let's look for headers.
        
        def get_points(sheet):
            points = {}
            # Looking for headers to find 'Name' and 'Implementation' columns
            name_col = 2 # Default B
            impl_col = 4 # Default D
            
            for row in range(1, 10):
                for col in range(1, 10):
                    val = str(sheet.cell(row, col).value).lower() if sheet.cell(row, col).value else ""
                    if 'name' in val: name_col = col
                    if 'implement' in val or 'support' in val: impl_col = col
            
            for row in range(name_col + 1, sheet.max_row + 1):
                name = sheet.cell(row, name_col).value
                if name and isinstance(name, str) and len(name) > 1:
                    impl = sheet.cell(row, impl_col).value
                    points[name] = str(impl).strip() if impl else "Unimplemented"
            return points

        pub_points = get_points(pub_sheet)
        eng_points = get_points(eng_sheet)
        
        all_point_names = sorted(list(set(pub_points.keys()) | set(eng_points.keys())))
        
        for p in all_point_names:
            p_val = pub_points.get(p, "N/A")
            e_val = eng_points.get(p, "N/A")
            
            if p_val != e_val:
                results["models"][model_id]["points"][p] = {
                    "public": p_val,
                    "engineering": e_val,
                    "change": "Implementation Status Mismatch"
                }

    return results

# File paths
public_pics = "/Users/davidhona/Downloads/UPDATED_FranklinWH_Modbus_PICS_SM-000028.xlsx"
engineering_pics = "/Users/davidhona/Downloads/PICS_span_20230711_SPANcomments20230803.xlsx"

audit_results = audit_pics(public_pics, engineering_pics)
print(json.dumps(audit_results, indent=2))
